#!/usr/bin/env python3
"""
Test del parser + detector + normalizador + matcher
contra los Excel reales de BetaDomino.
No escribe en base de datos.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl

ATTACHMENTS = Path("/home/workdir/attachments")
FILES = [
    ATTACHMENTS / "Resultado_Total_TORNEO_COPA_NATALE_BONGIOVANNI_Ronda_7_2026-09-22.xlsx",
    ATTACHMENTS / "Resultado_Total_3er_Torneo_de_Domino_El_Dividive_Pop_Club_Ronda_7_2026-09-22.xlsx",
]


def normalize_name(name: Optional[str]) -> str:
    if not name:
        return ""
    n = unicodedata.normalize("NFD", name)
    n = "".join(c for c in n if unicodedata.category(c) != "Mn")
    n = n.upper()
    n = re.sub(r"[^A-Z0-9\s]", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    return n


def normalize_id(val) -> Optional[str]:
    if val is None or val == "":
        return None
    s = str(val).strip()
    if not s:
        return None
    cleaned = s.lstrip("0") or "0"
    return cleaned


def normalize_club(club: Optional[str]) -> str:
    if not club:
        return ""
    n = normalize_name(club)
    aliases = {
        "S MENDOZA": "S MENDOZA",
        "S.MENDOZA": "S MENDOZA",
        "SMENDOZA": "S MENDOZA",
        "SANTA MENDOZA": "S MENDOZA",
        "EL DIVIDIVE": "DIVIDIVE",
        "DIVIDIVE": "DIVIDIVE",
        "VALERA": "VALERA",
        "BETIJOQUE": "BETIJOQUE",
        "TRUJILLO": "TRUJILLO",
        "KM 23": "KM 23",
        "KM23": "KM 23",
    }
    return aliases.get(n, n)


def parse_avg(val) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return float(val) / 100 if val > 1 else float(val)
    s = str(val).replace("%", "").strip().replace(",", ".")
    try:
        n = float(s)
        return n / 100 if n > 1 else n
    except ValueError:
        return 0.0


def parse_int(val, fallback=0) -> int:
    if val is None or val == "":
        return fallback
    try:
        return int(val)
    except (TypeError, ValueError):
        return fallback


def parse_filename_meta(filename: str):
    base = re.sub(r"\.xlsx?$", "", filename, flags=re.I)
    date_m = re.search(r"(\d{4}-\d{2}-\d{2})$", base)
    date = date_m.group(1) if date_m else None
    round_m = re.search(r"Ronda[_\s-]?(\d+)", base, re.I)
    round_number = int(round_m.group(1)) if round_m else None
    name_m = re.search(r"Resultado[_\s-]?Total[_\s-]+(.+?)[_\s-]+Ronda", base, re.I)
    if name_m:
        name = re.sub(r"_+", " ", name_m.group(1)).strip()
    else:
        name = re.sub(r"^Resultado[_\s-]?Total[_\s-]*", "", base, flags=re.I)
        name = re.sub(r"[_\s-]*Ronda[_\s-]?\d+.*", "", name, flags=re.I)
        name = re.sub(r"_+", " ", name).strip() or None
    return name, round_number, date


def build_external_key(typ: str, name: str, date: Optional[str], round_number: Optional[int]) -> str:
    name_part = normalize_name(name).replace(" ", "_") or "UNKNOWN"
    date_part = date or "unknown-date"
    round_part = str(round_number) if round_number is not None else "unknown-round"
    return f"betadomino|{typ}|{name_part}|{date_part}|{round_part}"


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


@dataclass
class RawRow:
    row_index: int
    pos: int
    name_raw: str
    beta_id: Optional[str]
    club_raw: str
    pj: int
    pg: int
    pp: int
    efe: int
    pf: int
    pc: int
    pm: int
    avg: float
    sanctions: Optional[int] = None
    pair_id: Optional[str] = None


def parse_individual_sheet(ws) -> list:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(h or "").upper().strip() for h in rows[0]]

    def col(name):
        for i, h in enumerate(headers):
            if h == name or name in h:
                return i
        return -1

    idx_pos = col("POS")
    idx_name = col("JUGADOR")
    idx_id = next((i for i, h in enumerate(headers) if "ID" in h and "USUARIO" in h), -1)
    idx_club = next((i for i, h in enumerate(headers) if "CLUB" in h or "ESTADO" in h), -1)
    idx_pj, idx_pg, idx_pp = col("PJ"), col("PG"), col("PP")
    idx_efe, idx_pf, idx_pc, idx_pm = col("EFE"), col("PF"), col("PC"), col("PM")
    idx_avg = col("AVG")
    idx_san = next((i for i, h in enumerate(headers) if "SANCION" in h), -1)

    result = []
    for i, row in enumerate(rows[1:], start=1):
        if not row or all(c is None or c == "" for c in row):
            continue
        name_raw = str(row[idx_name] or "").strip() if idx_name >= 0 else ""
        if not name_raw:
            continue
        result.append(RawRow(
            row_index=i,
            pos=parse_int(row[idx_pos] if idx_pos >= 0 else None, len(result) + 1),
            name_raw=name_raw,
            beta_id=normalize_id(row[idx_id] if idx_id >= 0 else None),
            club_raw=str(row[idx_club] or "").strip() if idx_club >= 0 else "",
            pj=parse_int(row[idx_pj] if idx_pj >= 0 else None),
            pg=parse_int(row[idx_pg] if idx_pg >= 0 else None),
            pp=parse_int(row[idx_pp] if idx_pp >= 0 else None),
            efe=parse_int(row[idx_efe] if idx_efe >= 0 else None),
            pf=parse_int(row[idx_pf] if idx_pf >= 0 else None),
            pc=parse_int(row[idx_pc] if idx_pc >= 0 else None),
            pm=parse_int(row[idx_pm] if idx_pm >= 0 else None),
            avg=parse_avg(row[idx_avg] if idx_avg >= 0 else None),
            sanctions=parse_int(row[idx_san]) if idx_san >= 0 and row[idx_san] not in (None, "") else None,
        ))
    return result


def parse_pairs_individual_sheet(ws) -> list:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(h or "").upper().strip() for h in rows[0]]

    def col(name):
        for i, h in enumerate(headers):
            if h == name or name in h:
                return i
        return -1

    idx_pos = col("POS")
    idx_name = col("JUGADOR")
    idx_id = next((i for i, h in enumerate(headers) if "ID" in h and "USUARIO" in h), -1)
    idx_pair = next((i for i, h in enumerate(headers) if "PAREJA" in h or "EQUIPO" in h), -1)
    idx_club = next((i for i, h in enumerate(headers) if "CLUB" in h or "ESTADO" in h), -1)
    idx_pj, idx_pg, idx_pp = col("PJ"), col("PG"), col("PP")
    idx_efe, idx_pf, idx_pc, idx_pm = col("EFE"), col("PF"), col("PC"), col("PM")
    idx_avg = col("AVG")
    idx_san = next((i for i, h in enumerate(headers) if "SANCION" in h), -1)

    result = []
    for i, row in enumerate(rows[1:], start=1):
        if not row or all(c is None or c == "" for c in row):
            continue
        name_raw = str(row[idx_name] or "").strip() if idx_name >= 0 else ""
        if not name_raw:
            continue
        result.append(RawRow(
            row_index=i,
            pos=parse_int(row[idx_pos] if idx_pos >= 0 else None, len(result) + 1),
            name_raw=name_raw,
            beta_id=normalize_id(row[idx_id] if idx_id >= 0 else None),
            club_raw=str(row[idx_club] or "").strip() if idx_club >= 0 else "",
            pj=parse_int(row[idx_pj] if idx_pj >= 0 else None),
            pg=parse_int(row[idx_pg] if idx_pg >= 0 else None),
            pp=parse_int(row[idx_pp] if idx_pp >= 0 else None),
            efe=parse_int(row[idx_efe] if idx_efe >= 0 else None),
            pf=parse_int(row[idx_pf] if idx_pf >= 0 else None),
            pc=parse_int(row[idx_pc] if idx_pc >= 0 else None),
            pm=parse_int(row[idx_pm] if idx_pm >= 0 else None),
            avg=parse_avg(row[idx_avg] if idx_avg >= 0 else None),
            sanctions=parse_int(row[idx_san]) if idx_san >= 0 and row[idx_san] not in (None, "") else None,
            pair_id=str(row[idx_pair]).strip() if idx_pair >= 0 and row[idx_pair] is not None else None,
        ))
    return result


def parse_file(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    upper = [s.upper() for s in sheet_names]

    name, round_number, date = parse_filename_meta(path.name)
    fhash = file_hash(path)

    is_pairs = any("RANKING" in s and ("PAREJA" in s or "EQUIPO" in s) for s in upper) and \
               any("JUGADOR" in s and "INDIVIDUAL" in s for s in upper)

    if is_pairs:
        detected = "parejas"
        sheet_name = next(s for s in sheet_names if "JUGADOR" in s.upper() and "INDIVIDUAL" in s.upper())
        rows = parse_pairs_individual_sheet(wb[sheet_name])
    else:
        detected = "individual"
        sheet_name = next((s for s in sheet_names if "RESULTADO" in s.upper()), sheet_names[0])
        rows = parse_individual_sheet(wb[sheet_name])

    tournament_name = name or ("Torneo Parejas" if detected == "parejas" else "Torneo Individual")
    external_key = build_external_key(detected, tournament_name, date, round_number)

    return {
        "file": path.name,
        "fileHash": fhash,
        "detectedType": detected,
        "tournament": {
            "name": tournament_name,
            "date": date,
            "roundNumber": round_number,
        },
        "externalKey": external_key,
        "rows": rows,
        "sheetNames": sheet_names,
    }


@dataclass
class ExistingPlayer:
    id: str
    display_name: str
    normalized_name: str
    club: Optional[str]
    normalized_club: Optional[str]
    beta_ids: list = field(default_factory=list)


def levenshtein(a: str, b: str) -> int:
    if a == b:
        return 0
    if not a:
        return len(b)
    if not b:
        return len(a)
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a):
        curr = [i + 1]
        for j, cb in enumerate(b):
            cost = 0 if ca == cb else 1
            curr.append(min(prev[j] + cost, prev[j + 1] + 1, curr[j] + 1))
        prev = curr
    return prev[-1]


def match_players(rows: list, existing: list):
    by_beta = defaultdict(list)
    by_name = defaultdict(list)
    by_name_club = defaultdict(list)

    for p in existing:
        for bid in p.beta_ids:
            k = normalize_id(bid)
            if k:
                by_beta[k].append(p)
        nn = p.normalized_name or normalize_name(p.display_name)
        if nn:
            by_name[nn].append(p)
            nc = p.normalized_club or normalize_club(p.club)
            by_name_club[f"{nn}|{nc}"].append(p)

    stats = dict(total=len(rows), exact_id=0, exact_name=0, fuzzy_name=0, new=0, conflict=0)
    results = []

    for row in rows:
        nn = normalize_name(row.name_raw)
        nc = normalize_club(row.club_raw)
        bid = normalize_id(row.beta_id)
        status = "new"
        confidence = 1.0
        reason = "Jugador no encontrado → crear nuevo"
        proposed = None

        if bid and bid in by_beta:
            cands = by_beta[bid]
            if len(cands) == 1:
                p = cands[0]
                pn = p.normalized_name or normalize_name(p.display_name)
                if pn and nn and levenshtein(pn, nn) > 3 and len(pn) >= 5 and len(nn) >= 5:
                    status, confidence = "conflict", 0.4
                    reason = f"Mismo betaId ({bid}) pero nombre distinto"
                    stats["conflict"] += 1
                else:
                    status, confidence = "exact_id", 1.0
                    reason = f"betaId exacto {bid}"
                    proposed = p
                    stats["exact_id"] += 1
            else:
                status, confidence = "conflict", 0.3
                reason = f"Varios jugadores con betaId {bid}"
                stats["conflict"] += 1
        elif nn and f"{nn}|{nc}" in by_name_club:
            cands = by_name_club[f"{nn}|{nc}"]
            if len(cands) == 1:
                status, confidence = "exact_name", 0.95
                reason = "Nombre + club exactos"
                proposed = cands[0]
                stats["exact_name"] += 1
            else:
                status, confidence = "conflict", 0.5
                reason = "Varios con mismo nombre+club"
                stats["conflict"] += 1
        elif nn and nn in by_name:
            cands = by_name[nn]
            if len(cands) == 1:
                status, confidence = "exact_name", 0.85
                reason = "Nombre único exacto"
                proposed = cands[0]
                stats["exact_name"] += 1
            else:
                status, confidence = "conflict", 0.5
                reason = "Varios con mismo nombre"
                stats["conflict"] += 1
        elif len(nn) >= 8:
            fuzzy = []
            for p in existing:
                pn = p.normalized_name or normalize_name(p.display_name)
                if not pn or len(pn) < 6:
                    continue
                d = levenshtein(nn, pn)
                if d <= 2:
                    fuzzy.append((p, d))
            if len(fuzzy) == 1:
                p, d = fuzzy[0]
                status = "fuzzy_name"
                confidence = 0.8 if d == 0 else (0.75 if d == 1 else 0.7)
                reason = f"Fuzzy distancia {d}"
                proposed = p
                stats["fuzzy_name"] += 1
            elif len(fuzzy) > 1:
                status, confidence = "conflict", 0.4
                reason = "Varios candidatos fuzzy"
                stats["conflict"] += 1
            else:
                stats["new"] += 1
        else:
            stats["new"] += 1

        results.append({
            "pos": row.pos,
            "name": row.name_raw,
            "betaId": row.beta_id,
            "club": row.club_raw,
            "status": status,
            "confidence": confidence,
            "reason": reason,
            "proposed": proposed.display_name if proposed else None,
        })

    return results, stats


def main():
    print("=" * 70)
    print("TEST PARSER + MATCHER — BetaDomino Excel reales")
    print("=" * 70)

    sample_players = [
        ExistingPlayer("p1", "LUIS SERRADA", normalize_name("LUIS SERRADA"), "VALERA", normalize_club("VALERA"), ["6"]),
        ExistingPlayer("p2", "JULCOR PAREDES", normalize_name("JULCOR PAREDES"), "VALERA", normalize_club("VALERA"), ["3"]),
        ExistingPlayer("p3", "Mariela de Calomino", normalize_name("Mariela de Calomino"), "S.Mendoza", normalize_club("S.Mendoza"), ["23"]),
        ExistingPlayer("p4", "DOUGLAS PERDOMO", normalize_name("DOUGLAS PERDOMO"), "DIVIDIVE", normalize_club("DIVIDIVE"), ["1", "27"]),
    ]

    for path in FILES:
        if not path.exists():
            print(f"\n[SKIP] No existe: {path}")
            continue

        print(f"\n{'─' * 70}")
        print(f"ARCHIVO: {path.name}")
        print(f"{'─' * 70}")

        result = parse_file(path)
        print(f"  Tipo detectado : {result['detectedType']}")
        print(f"  Torneo         : {result['tournament']['name']}")
        print(f"  Fecha          : {result['tournament']['date']}")
        print(f"  Ronda          : {result['tournament']['roundNumber']}")
        print(f"  externalKey    : {result['externalKey']}")
        print(f"  fileHash       : {result['fileHash'][:16]}…")
        print(f"  Hojas          : {result['sheetNames']}")
        print(f"  Filas parseadas: {len(result['rows'])}")

        if result["rows"]:
            print("\n  Primeras 5 filas:")
            for r in result["rows"][:5]:
                print(f"    POS {r.pos:>2} | {r.name_raw:<25} | ID={r.beta_id or '-':<6} | {r.club_raw:<12} | "
                      f"PJ={r.pj} PG={r.pg} PP={r.pp} EFE={r.efe} AVG={r.avg:.0%}")

        _, stats_empty = match_players(result["rows"], [])
        print(f"\n  Matching (DB vacía) → new={stats_empty['new']}  (todo nuevo, esperado)")

        matched, stats = match_players(result["rows"], sample_players)
        print(f"  Matching (con muestra) → "
              f"exact_id={stats['exact_id']}  exact_name={stats['exact_name']}  "
              f"fuzzy={stats['fuzzy_name']}  new={stats['new']}  conflict={stats['conflict']}")

        hits = [m for m in matched if m["status"] != "new"]
        if hits:
            print("\n  Coincidencias encontradas:")
            for m in hits[:10]:
                print(f"    [{m['status']:11}] {m['name']:<25} → {m['proposed'] or m['reason']}")

    print(f"\n{'=' * 70}")
    print("FIN DEL TEST")
    print("=" * 70)


if __name__ == "__main__":
    main()
